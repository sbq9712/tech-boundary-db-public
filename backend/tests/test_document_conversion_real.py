"""Opt-in real Docling/OCR/model tests. Generates synthetic files, no private data.
RUN_DOCLING_REAL_TESTS=1 pytest -s tests/test_document_conversion_real.py
"""

import json
import os
import uuid
from pathlib import Path

import pytest

pytestmark = pytest.mark.skipif(
    os.getenv("RUN_DOCLING_REAL_TESTS") != "1",
    reason="real Docling models: set RUN_DOCLING_REAL_TESTS=1",
)


def make_documents(root):
    from docx import Document
    from PIL import Image, ImageDraw, ImageFont
    from pptx import Presentation
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    text = (
        "Synthetic battery evidence: capacity retention is 95 percent after 100 cycles."
    )
    (root / "plain.txt").write_text("# Synthetic evidence\n\n" + text, encoding="utf-8")
    doc = Document()
    doc.add_heading("Synthetic evidence", level=1)
    doc.add_paragraph(text)
    doc.save(root / "report.docx")
    table_doc = Document()
    table_doc.add_heading("Synthetic results", level=1)
    table = table_doc.add_table(rows=3, cols=2)
    for row, values in zip(
        table.rows, [("Metric", "Value"), ("Capacity", "95 percent"), ("Cycles", "100")]
    ):
        for cell, value in zip(row.cells, values):
            cell.text = value
    table_doc.save(root / "table.docx")
    pres = Presentation()
    slide = pres.slides.add_slide(pres.slide_layouts[1])
    slide.shapes.title.text = "Synthetic evidence"
    slide.placeholders[1].text = text
    pres.save(root / "slides.pptx")
    pdf = canvas.Canvas(str(root / "digital.pdf"))
    pdf.setFont("Helvetica-Bold", 18)
    pdf.drawString(50, 780, "Synthetic evidence")
    pdf.setFont("Helvetica", 11)
    pdf.drawString(50, 740, text)
    pdf.save()
    image = Image.new("RGB", (1600, 700), "white")
    draw = ImageDraw.Draw(image)
    font_path = os.getenv(
        "TEST_OCR_FONT", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    )
    font = ImageFont.truetype(font_path, 36)
    draw.text((50, 100), "SCANNED SYNTHETIC EVIDENCE", font=font, fill="black")
    draw.text(
        (50, 200), "Battery capacity retention: 95 percent", font=font, fill="black"
    )
    draw.text((50, 300), "Test duration: 100 cycles", font=font, fill="black")
    scan = canvas.Canvas(str(root / "scan.pdf"), pagesize=(800, 350))
    scan.drawImage(ImageReader(image), 0, 0, width=800, height=350)
    scan.save()  # image only: no text drawing, no text layer
    return [
        root / name
        for name in [
            "plain.txt",
            "report.docx",
            "slides.pptx",
            "table.docx",
            "digital.pdf",
            "scan.pdf",
        ]
    ]


@pytest.fixture(autouse=True)
def optional_fixture_tokenizer(monkeypatch):
    # Only for isolated conversion tests when the real tokenizer cannot be downloaded.
    # The default path always uses the configured bge-m3 tokenizer.
    if os.getenv("TEST_DOCUMENT_TOKENIZER") != "fixture":
        return
    from app.document_ingestion import chunker
    from docling.chunking import HybridChunker
    from docling_core.transforms.chunker.tokenizer.base import BaseTokenizer

    class FixtureTokenizer(BaseTokenizer):
        def count_tokens(self, text):
            return len(text.split())

        def get_max_tokens(self):
            return 512

        def get_tokenizer(self):
            return self.count_tokens

    engine = HybridChunker(tokenizer=FixtureTokenizer())
    monkeypatch.setattr(chunker, "chunker", lambda: engine)


@pytest.fixture(scope="module")
def synthetic_documents(tmp_path_factory):
    root = tmp_path_factory.mktemp("real-docling")
    return root, {p.name: p for p in make_documents(root)}


def test_isolated_pdf_worker(synthetic_documents, monkeypatch):
    monkeypatch.delenv("TECHDB_DOCLING_FORMULAS", raising=False)
    if os.getenv("TEST_DOCUMENT_TOKENIZER") == "fixture":
        pytest.skip("child worker requires the real tokenizer")
    from app.document_ingestion.runner import ingest
    from app.document_ingestion.storage import sha256

    root, files = synthetic_documents
    original = files["digital.pdf"]
    assets, markdown, chunks = ingest(
        root=root, original=original, document_id=str(uuid.uuid4()),
        filename=original.name, checksum=sha256(original),
    )
    assert assets["pipeline"]["formulas"] == "false"
    assert "Synthetic" in markdown
    assert chunks and (root / assets["bundle"] / "document.json").is_file()


@pytest.mark.parametrize(
    "filename",
    [
        "plain.txt",
        "report.docx",
        "slides.pptx",
        "table.docx",
        "digital.pdf",
        "scan.pdf",
    ],
)
def test_real_conversion_assets_and_replay(synthetic_documents, filename):
    from app.document_ingestion.service import ingest
    from app.document_ingestion.storage import sha256

    root, files = synthetic_documents
    original = files[filename]
    document_id = str(uuid.uuid5(uuid.NAMESPACE_URL, "techdb-synthetic/" + filename))
    assets, markdown, chunks = ingest(
        root=root,
        original=original,
        document_id=document_id,
        filename=filename,
        checksum=sha256(original),
    )
    directory = root / assets["bundle"]
    assert original.is_file() and markdown and chunks
    for name in [
        "document.md",
        "document.json",
        "chunks.jsonl",
        "chunk_metadata.jsonl",
        "conversion.json",
    ]:
        assert (directory / name).is_file()
    document_json = json.loads(
        (directory / "document.json").read_text(encoding="utf-8")
    )
    assert "body" in document_json and "texts" in document_json
    assert all(c["metadata"]["source_file"] == filename for c in chunks)
    assert all(c["metadata"]["document_id"] == document_id for c in chunks)
    if filename.endswith(".pdf"):
        assert any(c["metadata"]["page_number"] == 1 for c in chunks)
    if filename == "scan.pdf":
        assert assets["conversion"]["ocr_triggered"] is True
        assert "95" in markdown and "100" in markdown
    if filename == "table.docx":
        assert document_json["tables"]
        assert "Capacity" in markdown and "95" in markdown
        assert any(
            "#/tables/" in json.dumps(c["metadata"]["provenance"]) for c in chunks
        )
    replay = ingest(
        root=root,
        original=original,
        document_id=document_id,
        filename=filename,
        checksum=sha256(original),
        previous=assets,
    )
    assert replay == (assets, markdown, chunks)
    report = {
        "tokenizer": os.getenv("TEST_DOCUMENT_TOKENIZER", "configured bge-m3"),
        "document_id": document_id,
        "file_type": assets["conversion"]["input_format"],
        "filename": filename,
        "chunks_count": len(chunks),
        "ocr_triggered": assets["conversion"]["ocr_triggered"],
        "markdown_example": markdown[:500],
        "chunk_metadata_example": chunks[0]["metadata"],
    }
    print(json.dumps(report, ensure_ascii=False, indent=2))
    destination = os.getenv("DOCLING_TEST_REPORT_DIR")
    if destination:
        Path(destination).mkdir(parents=True, exist_ok=True)
        (Path(destination) / (filename + ".json")).write_text(
            json.dumps(report, ensure_ascii=False, indent=2)
        )
